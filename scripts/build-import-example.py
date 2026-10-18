"""Build the static downloadable import example; speech is supplied as a reviewed local asset."""

import argparse
import io
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from PIL import Image, ImageDraw

from juya_api.media_api import decoded_audio


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--audio-path", required=True, type=Path)
    arguments = parser.parse_args()
    audio = arguments.audio_path.read_bytes()
    content_type, duration = decoded_audio(audio)
    if content_type != "audio/wav" or duration <= 0:
        raise ValueError("Reviewed WAV required")
    document = Document()
    document.core_properties.author = "句芽英语"
    document.core_properties.last_modified_by = "句芽英语"
    document.add_heading("在咖啡店点单", level=1)
    document.add_paragraph("英文标题：Ordering Coffee")
    document.add_heading("场景对话", level=2)
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["人物", "英文", "中文"], strict=True):
        cell.text = value
    for cell, value in zip(
        table.add_row().cells,
        ["Ivy", "Could I get a latte, please?", "请给我一杯拿铁。"],
        strict=True,
    ):
        cell.text = value
    document.add_heading("重点词汇", level=2)
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(
        table.rows[0].cells,
        ["词汇", "本句释义（选填）", "音标（选填）", "对应英文原句（建议填写）"],
        strict=True,
    ):
        cell.text = value
    for cell, value in zip(
        table.add_row().cells,
        ["latte", "拿铁咖啡", "", "Could I get a latte, please?"],
        strict=True,
    ):
        cell.text = value
    docx = io.BytesIO()
    document.save(docx)
    image = Image.new("RGB", (1200, 900), "#f1f3e8")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(
        (225, 300, 830, 650), radius=50, fill="#fffdf7", outline="#347553", width=12
    )
    draw.arc((680, 340, 990, 590), -90, 90, fill="#347553", width=18)
    draw.ellipse((225, 260, 830, 380), fill="#d7bb97", outline="#347553", width=12)
    draw.ellipse((170, 665, 950, 730), fill="#dce5d5")
    for x in [420, 520, 620]:
        draw.arc((x, 140, x + 90, 255), 80, 270, fill="#93ad94", width=8)
    picture = io.BytesIO()
    image.save(picture, "PNG")
    target = Path(__file__).resolve().parents[1] / "src/juya_api/assets/import-example.zip"
    target.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        archive.writestr("在咖啡店点单/内容.docx", docx.getvalue())
        archive.writestr("在咖啡店点单/场景图.png", picture.getvalue())
        archive.writestr("在咖啡店点单/完整音频.wav", audio)
        archive.writestr(
            "填写说明.txt",
            "一张卡片一个文件夹，配套一份内容DOCX、一张完整场景图及一份完整对话音频。\n"
            "示例可替换为你自己的资料；导入只生成草稿，发布前仍需核对。\n"
            "独立DOCX入口允许多个一级标题，ZIP的每个文件夹只能包含一个一级标题。\n"
            "一级标题填写中文标题，紧接正文填写英文标题：Ordering Coffee；新发布需双标题完整。\n"
            "音频支持MP3、M4A、WAV、AAC；保留原件并兼容处理为MP3。\n"
            "本句释义和音标独立选填；语块选填，不需要上传词汇音频或语块图片。\n",
        )
    print(f"Built {target.name}; {duration:.2f}s reviewed sample audio")


if __name__ == "__main__":
    main()
