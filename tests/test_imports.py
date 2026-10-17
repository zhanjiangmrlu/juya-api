"""Fixed DOCX parsing and hostile ZIP input never use filesystem extraction."""

import io
import stat
import zipfile

import pytest
from docx import Document


def document(*names):
    doc = Document()
    for name in names:
        doc.add_heading(name, level=1)
        doc.add_heading("场景对话", level=2)
        table = doc.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells, ["人物", "英文", "中文"], strict=True):
            cell.text = text
        for cell, text in zip(
            table.add_row().cells,
            ["Ivy", "Could I get a latte, please?", "请给我一杯拿铁。"],
            strict=True,
        ):
            cell.text = text
        doc.add_heading("重点词汇", level=2)
        table = doc.add_table(rows=1, cols=4)
        for cell, text in zip(
            table.rows[0].cells,
            ["词汇", "本句释义", "音标（选填）", "对应英文原句（建议填写）"],
            strict=True,
        ):
            cell.text = text
        for cell, text in zip(
            table.add_row().cells,
            ["latte", "拿铁咖啡", "", "Could I get a latte, please?"],
            strict=True,
        ):
            cell.text = text
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def archive(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as data:
        for name, content in entries:
            data.writestr(name, content)
    return output.getvalue()


def parser():
    try:
        from juya_api.import_parser import checked_archive, parse_docx
    except ImportError as error:
        pytest.fail(f"Fixed safe import parsing missing: {error}")
    return parse_docx, checked_archive


def test_docx_uses_explicit_h1_boundaries_and_preserves_dialogue_and_confirmed_source():
    parse, _ = parser()
    values = parse(document("咖啡店", "机场"))
    assert [v["name"] for v in values] == ["咖啡店", "机场"]
    assert values[0]["dialogue"][0]["chinese"] == "请给我一杯拿铁。"
    word = values[0]["vocabulary"][0]
    assert word["meaning"] == "拿铁咖啡"
    assert word["source"]["quote"] == "latte"
    assert values[0]["chunks"] == []
    assert values[0]["dialogue"][0]["id"] != values[1]["dialogue"][0]["id"]
    doc = Document()
    doc.add_paragraph("卡片名称，任意空行不能当分篇标题")
    stream = io.BytesIO()
    doc.save(stream)
    from juya_api.errors import DomainError

    with pytest.raises(DomainError, match="docx_structure"):
        parse(stream.getvalue())


@pytest.mark.parametrize(
    "path",
    [
        "../escape.docx",
        "/absolute.docx",
        "C:/bad.docx",
        "a/../../bad.docx",
        "a\\bad.docx",
        "a//bad.docx",
    ],
)
def test_zip_rejects_unsafe_names_without_extracting(path):
    _, checked = parser()
    from juya_api.errors import DomainError

    with pytest.raises(DomainError, match="archive_unsafe"):
        checked(archive([(path, b"content")]))


def test_zip_rejects_duplicate_paths_symlinks_and_expansion_bombs():
    _, checked = parser()
    from juya_api.errors import DomainError

    with pytest.raises(DomainError, match="archive_unsafe"):
        checked(archive([("one/doc.docx", b"one"), ("one/doc.docx", b"two")]))
    link = zipfile.ZipInfo("one/link")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(DomainError, match="archive_unsafe"):
        checked(archive([(link, b"../../elsewhere")]))
    with pytest.raises(DomainError, match="archive_limit"):
        checked(archive([("one/padding", b"x" * 2_000_000)]), expanded_limit=1_000_000)


def test_docx_external_entities_and_unreadable_documents_are_rejected():
    parse, _ = parser()
    from juya_api.errors import DomainError

    xml = b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&e;</w:t></w:r></w:p></w:body></w:document>'
    with pytest.raises(DomainError, match="docx_invalid"):
        parse(archive([("word/document.xml", xml)]))
    with pytest.raises(DomainError, match="docx_invalid"):
        parse(b"not a zip file")
