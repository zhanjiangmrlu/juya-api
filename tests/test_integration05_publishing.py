"""Explicit publication and canonical document ownership never depend on viewing a preview."""

import io
import zipfile

from docx import Document

from tests.content_helpers import key, recording
from tests.test_integration05_catalog import catalog_topic
from tests.test_listening import A, M, article, policy, post
from tests.test_listening_import_updates import bind
from tests.test_listening_imports import assign, document, package, upload


def test_downloaded_template_parses_with_the_same_canonical_fields(admin):
    from juya_api.listening_imports import parse_docx

    response = admin.get(A + "/template")
    assert response.status_code == 200
    assert "wordprocessingml.document" in response.headers["content-type"]
    payload, hints = parse_docx(response.content)
    assert payload["title_zh"] and payload["title_en"] and payload["body_en"]
    assert payload["translation_zh"] and len(payload["expressions"]) == 5
    assert hints["主题"] and hints["分类"]


def test_single_publish_needs_no_manual_preview_and_replays_one_version(admin):
    _, topic = catalog_topic(admin)
    row = article(admin, topic["id"])
    headers = key()
    path = A + f"/articles/{row['id']}/publish"
    response = admin.post(path, headers=headers, json={"revision": row["revision"]})
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "published"
    assert (
        admin.post(path, headers=headers, json={"revision": row["revision"]}).json()
        == response.json()
    )


def test_batch_publish_keeps_valid_success_when_another_draft_is_incomplete(admin):
    _, topic = catalog_topic(admin)
    valid = article(admin, topic["id"])
    invalid = post(admin, A + "/articles", {"topic_id": topic["id"], "payload": {}})
    response = post(
        admin,
        A + "/articles/batch-publish",
        {"items": [{"id": r["id"], "revision": r["revision"]} for r in [valid, invalid]]},
    )
    assert [item["status"] for item in response["items"]] == ["success", "failed"]
    assert admin.get(A + f"/articles/{invalid['id']}").json()["status"] == "draft"


def test_template_alias_and_english_ownership_are_matched_without_creating_catalog(admin):
    category, topic = catalog_topic(admin)
    response = admin.put(
        A + f"/categories/{category['id']}",
        headers=key(),
        json={"name": category["name"], "name_en": "Everyday", "revision": category["revision"]},
    )
    assert response.status_code == 200
    doc = Document(io.BytesIO(document()))
    doc.paragraphs[0].text, doc.paragraphs[0].style = "午间休息", "Heading 1"
    doc.paragraphs[2].insert_paragraph_before("分类: Everyday")
    doc.paragraphs[3].insert_paragraph_before("所属主题：A short break")
    content = io.BytesIO()
    doc.save(content)
    package_stream = io.BytesIO()
    with zipfile.ZipFile(package_stream, "w") as archive:
        archive.writestr("one/content.docx", content.getvalue())
        archive.writestr("one/complete.wav", recording())
    batch = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package_stream.getvalue())
    candidate = batch["candidates"][0]
    assert candidate["state"] == "complete", candidate["errors"]
    assert candidate["topic_id"] == topic["id"]
    assert len(admin.get(A + "/topics").json()["items"]) == 1


def test_import_publish_is_atomic_per_candidate_and_recoverable(admin, mini):
    _, topic = catalog_topic(admin)
    batch = upload(
        admin, post(admin, A + "/imports", {"mode": "zip"}), package("完整"), package("缺项")
    )
    for candidate in batch["candidates"]:
        batch = assign(admin, batch, candidate["id"], topic["id"])
    good, bad = batch["candidates"]
    batch = admin.put(
        A + f"/imports/{batch['id']}/candidates/{bad['id']}",
        headers=key(),
        json={
            "revision": batch["revision"],
            "topic_id": topic["id"],
            "payload": {**bad["payload"], "title_en": ""},
        },
    ).json()
    preview = post(
        admin,
        A + f"/imports/{batch['id']}/preview",
        {
            "revision": batch["revision"],
            "candidate_ids": [good["id"], bad["id"]],
            "submit_action": "publish",
        },
    )
    headers = key()
    body = {"preview_token": preview["preview_token"], "submit_action": "publish"}
    response = admin.post(A + f"/imports/{batch['id']}/save", headers=headers, json=body)
    assert response.status_code == 200, response.text
    result = response.json()
    assert [item["status"] for item in result["items"]] == ["success", "failed"]
    assert result["items"][0]["published_version_id"]
    assert admin.get(A + "/articles").json()["total"] == 1
    assert result["batch"]["mode"] == "zip"
    assert (
        admin.post(A + f"/imports/{batch['id']}/save", headers=headers, json=body).json() == result
    )
    policy(admin, 1)
    assert mini.get(M + f"/articles/{result['items'][0]['article_id']}").status_code == 200


def test_publish_explicit_update_of_unchanged_draft_does_not_skip_unpublished_changes(admin):
    _, topic = catalog_topic(admin)
    row = article(admin, topic["id"])
    live = post(admin, A + f"/articles/{row['id']}/publish", {"revision": row["revision"]})
    changed = admin.put(
        A + f"/articles/{live['id']}",
        headers=key(),
        json={
            "revision": live["revision"],
            "topic_id": topic["id"],
            "payload": {**live["payload"], "title_zh": "未发布修改"},
        },
    ).json()
    batch = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package("用于明确更新"))
    cid = batch["candidates"][0]["id"]
    batch = bind(admin, batch, cid, changed, payload=changed["payload"])
    preview = post(
        admin,
        A + f"/imports/{batch['id']}/preview",
        {
            "revision": batch["revision"],
            "candidate_ids": [cid],
            "submit_action": "publish",
        },
    )
    saved = post(
        admin,
        A + f"/imports/{batch['id']}/save",
        {
            "preview_token": preview["preview_token"],
            "submit_action": "publish",
        },
    )
    assert saved["items"][0]["status"] == "success"
    current = admin.get(A + f"/articles/{live['id']}/published").json()
    assert current["payload"]["title_zh"] == "未发布修改"
