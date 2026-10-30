"""The downloadable sample uses exactly the fields accepted by the listening parser."""

import io

from docx import Document


def template_document() -> bytes:
    """Build one H1 article, explicit ownership metadata and five complete expression rows."""
    document = Document()
    document.add_heading("午间休息", level=1)
    document.add_paragraph("英文标题：A Short Break")
    document.add_paragraph("分类：请填写已有分类的中文或英文名称")
    document.add_paragraph("所属主题：请填写已有主题的英文或中文名称")
    document.add_heading("英文正文", level=2)
    document.add_paragraph(
        "Take a break and drink water. Slow down, take a breath, and feel better."
    )
    document.add_heading("全文中文译文", level=2)
    document.add_paragraph("休息一下，喝点水。放慢脚步，深呼吸，感觉更好一些。")
    document.add_heading("5组语块", level=2)
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["英文", "中文释义", "替换例句"], strict=True):
        cell.text = value
    for values in [
        ["take a break", "休息一下", "Let's take a break."],
        ["drink water", "喝水", "Remember to drink water."],
        ["slow down", "放慢脚步", "We can slow down."],
        ["take a breath", "深呼吸", "Stop and take a breath."],
        ["feel better", "感觉好些", "I feel better now."],
    ]:
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
