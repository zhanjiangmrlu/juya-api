"""Real DOCX/ZIP import staging, deduplication, explicit saving and recovery."""

import io
import zipfile

import pytest
from docx import Document

from tests.content_helpers import key, recording
from tests.test_listening import A, post, topic


def document(title="午间休息"):
    doc = Document()
    doc.add_paragraph("中文标题：" + title)
    doc.add_paragraph("英文标题：A short break")
    doc.add_paragraph("英文正文")
    doc.add_paragraph("We rest.")
    doc.add_paragraph("We drink water.")
    doc.add_paragraph("全文中文译文")
    doc.add_paragraph("我们休息。\n我们喝水。")
    doc.add_paragraph("5组语块")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["英文", "中文释义", "替换例句"], strict=True):
        cell.text = text
    for _ in range(5):
        for cell, text in zip(
            table.add_row().cells, ["take a break", "休息", "Take a break."], strict=True
        ):
            cell.text = text
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def package(title="午间休息", duplicate=False):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for folder in ["one", "two"] if duplicate else ["one"]:
            archive.writestr(folder + "/content.docx", document(title))
            archive.writestr(folder + "/complete.wav", recording())
    return output.getvalue()


def upload(admin, batch, *packages):
    result = admin.post(
        A + f"/imports/{batch['id']}/files?revision={batch['revision']}",
        headers=key(),
        files=[("files", ("same.zip", p, "application/zip")) for p in packages],
    )
    assert result.status_code == 200, result.text
    return result.json()


def assign(admin, batch, candidate_id, topic_id):
    candidate = next(c for c in batch["candidates"] if c["id"] == candidate_id)
    result = admin.put(
        A + f"/imports/{batch['id']}/candidates/{candidate_id}",
        headers=key(),
        json={"revision": batch["revision"], "topic_id": topic_id, "payload": candidate["payload"]},
    )
    assert result.status_code == 200, result.text
    return result.json()


def test_zip_docx_natural_paragraphs_duplicate_skip_and_explicit_save(admin, system):
    t = topic(admin)
    batch = post(admin, A + "/imports", {"mode": "zip"})
    batch = upload(admin, batch, package(duplicate=True))
    assert batch["counts"]["duplicate"] == 1
    candidate = next(c for c in batch["candidates"] if c["state"] != "duplicate")
    assert candidate["payload"]["body_en"] == "We rest.\nWe drink water."
    assert len(candidate["payload"]["expressions"]) == 5
    assert admin.get(A + "/articles").json()["total"] == 0
    batch = assign(admin, batch, candidate["id"], t["id"])
    preview = post(
        admin,
        A + f"/imports/{batch['id']}/preview",
        {"revision": batch["revision"], "candidate_ids": [candidate["id"]]},
    )
    result = post(
        admin, A + f"/imports/{batch['id']}/save", {"preview_token": preview["preview_token"]}
    )
    assert result["items"][0]["status"] == "success"
    assert admin.get(A + "/articles").json()["total"] == 1
    from sqlalchemy import func, select

    from juya_api.listening_models import ListeningJob

    with system["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(ListeningJob)) == 1
    repeated = post(admin, A + "/imports", {"mode": "zip"})
    repeated = upload(admin, repeated, package())
    assert repeated["counts"]["duplicate"] == 1


def test_candidate_remove_is_durable_and_stales_confirmation(admin):
    t = topic(admin)
    batch = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package())
    candidate_id = batch["candidates"][0]["id"]
    batch = assign(admin, batch, candidate_id, t["id"])
    preview = post(
        admin,
        A + f"/imports/{batch['id']}/preview",
        {"revision": batch["revision"], "candidate_ids": [candidate_id]},
    )
    batch = post(
        admin,
        A + f"/imports/{batch['id']}/remove",
        {"revision": batch["revision"], "candidate_ids": [candidate_id]},
    )
    assert batch["counts"]["savable"] == 0
    assert admin.get(A + f"/imports/{batch['id']}").json()["candidates"][0]["removed"]
    result = admin.post(
        A + f"/imports/{batch['id']}/save",
        headers=key(),
        json={"preview_token": preview["preview_token"]},
    )
    assert result.status_code == 409
    assert admin.get(A + "/articles").json()["total"] == 0


def test_same_name_different_content_and_multi_zip_names_do_not_mix(admin):
    batch = upload(
        admin, post(admin, A + "/imports", {"mode": "zip"}), package("甲"), package("乙")
    )
    assert batch["counts"]["duplicate"] == 0
    assert len({c["package_id"] for c in batch["candidates"]}) == 2
    assert {c["payload"]["title_zh"] for c in batch["candidates"]} == {"甲", "乙"}


def test_unreadable_can_remove_but_cannot_save(admin):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("bad/content.docx", b"invalid")
    batch = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), output.getvalue())
    assert batch["candidates"][0]["state"] == "unreadable"
    cid = batch["candidates"][0]["id"]
    assert (
        admin.post(
            A + f"/imports/{batch['id']}/preview",
            headers=key(),
            json={"revision": batch["revision"], "candidate_ids": [cid]},
        ).status_code
        == 422
    )
    batch = post(
        admin,
        A + f"/imports/{batch['id']}/remove",
        {"revision": batch["revision"], "candidate_ids": [cid]},
    )
    assert batch["counts"]["removed"] == 1


def test_docx_preview_has_no_article_write_and_staged_audio_is_private(admin, mini):
    preview = admin.post(
        A + "/docx-preview",
        headers=key(),
        files={
            "file": (
                "article.docx",
                document(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["payload"]["title_en"] == "A short break"
    assert admin.get(A + "/articles").json()["total"] == 0
    batch = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package())
    candidate = batch["candidates"][0]
    medium = candidate["payload"]["audio_id"]
    path = A + f"/imports/{batch['id']}/media/{medium}"
    assert admin.get(path).status_code == 200
    assert mini.get(path).status_code == 404
    changed = post(
        admin,
        A + f"/imports/{batch['id']}/candidates/{candidate['id']}/timeline/review",
        {"revision": batch["revision"]},
    )
    assert changed["reviewed_at"] is not None
    assert admin.get(A + "/articles").json()["total"] == 0


@pytest.mark.parametrize("reupload", [False, True])
def test_deleted_article_no_longer_blocks_reimport(admin, reupload):
    """Deleted articles and their saved import candidates cannot block a new draft."""
    t = topic(admin)
    original = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package())
    cid = original["candidates"][0]["id"]
    original = assign(admin, original, cid, t["id"])
    preview = post(
        admin,
        A + f"/imports/{original['id']}/preview",
        {"revision": original["revision"], "candidate_ids": [cid]},
    )
    saved = post(
        admin,
        A + f"/imports/{original['id']}/save",
        {"preview_token": preview["preview_token"]},
    )
    old_id = saved["items"][0]["article_id"]
    old = admin.get(A + f"/articles/{old_id}").json()
    duplicate = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package())
    assert duplicate["counts"]["duplicate"] == 1
    post(admin, A + f"/articles/{old_id}/delete", {"revision": old["revision"]})
    if reupload:
        batch = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package())
    else:
        batch = admin.get(A + f"/imports/{duplicate['id']}").json()
        assert batch["revision"] > duplicate["revision"]
    candidate = batch["candidates"][0]
    assert candidate["state"] == "incomplete"
    assert candidate["duplicate_of"] is None
    assert batch["counts"]["savable"] == 1
    assert candidate["timeline"]["state"] == "pending"
    assert candidate["timeline"]["job_id"] is not None
    stable = admin.get(A + f"/imports/{batch['id']}").json()
    assert stable["revision"] == batch["revision"]
    assert stable["candidates"][0]["state"] == "incomplete"
    batch = assign(admin, batch, candidate["id"], t["id"])
    preview = post(
        admin,
        A + f"/imports/{batch['id']}/preview",
        {"revision": batch["revision"], "candidate_ids": [candidate["id"]]},
    )
    current = post(
        admin,
        A + f"/imports/{batch['id']}/save",
        {"preview_token": preview["preview_token"]},
    )
    assert current["items"][0]["status"] == "success"
    assert current["items"][0]["article_id"] != old_id
    assert admin.get(A + f"/articles/{old_id}").json()["deleted_at"] is not None
    assert admin.get(A + "/articles").json()["total"] == 1
    repeated = upload(admin, post(admin, A + "/imports", {"mode": "zip"}), package())
    assert repeated["counts"]["duplicate"] == 1
    assert repeated["candidates"][0]["duplicate_of"] == current["items"][0]["article_id"]
